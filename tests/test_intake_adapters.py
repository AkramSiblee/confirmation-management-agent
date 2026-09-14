"""
Tests for the intake-adapter mechanism (src/confirmation_agent/intake_adapters/).

Builds a small synthetic workbook that deliberately deviates from
docs/io-spec.md's canonical layout the same way example_client.py's
docstring describes (renamed Bank tab, shifted header row, renamed
columns, renamed Engagement Setup field) and asserts:

  1. Loading it with no adapter fails the way real tab/column drift
     fails today — loudly, not with silently wrong data.
  2. Loading it with adapter="example_client" produces the same shape
     and values load_all() would produce from a canonical workbook,
     proving the adapter normalizes the client's layout without
     touching schemas.py/io-spec.md or intake.py's validation logic.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from confirmation_agent import config, schemas
from confirmation_agent.intake import load_all
from confirmation_agent.intake_adapters import get_adapter


def _write_table(ws, header_excel_row: int, columns: list[str], rows: list[list]) -> None:
    for r in range(1, header_excel_row):
        ws.cell(row=r, column=1, value=f"(filler row {r})")
    for c, col in enumerate(columns, start=1):
        ws.cell(row=header_excel_row, column=c, value=col)
    for i, row_vals in enumerate(rows, start=1):
        for c, val in enumerate(row_vals, start=1):
            ws.cell(row=header_excel_row + i, column=c, value=val)


def _build_client_workbook(tmp_path):
    """A workbook shaped exactly like example_client.py's docstring describes:
    Bank tab renamed + shifted header + renamed columns, Engagement Setup
    field renamed. Every other tab already matches io-spec.md canonically."""
    wb = Workbook()
    wb.remove(wb.active)

    # Engagement Setup: canonical header row/usecols, but "Client" instead
    # of "Client Name" (engagement_field_map is what fixes this).
    ws = wb.create_sheet(config.SHEET_ENGAGEMENT_SETUP)
    header_excel_row = config.ENGAGEMENT_SETUP_HEADER_ROW + 1
    for r in range(1, header_excel_row):
        ws.cell(row=r, column=2, value=f"(filler row {r})")
    ws.cell(row=header_excel_row, column=2, value="Field")
    ws.cell(row=header_excel_row, column=3, value="Value")
    fields = [
        ("Client", "Acme Test Co"),
        ("Engagement Code", "ACM-FY26-001"),
        ("Period-End Date", "2026-12-31"),
        ("Group Reporting Currency", "USD"),
        ("Group Materiality", 100000),
        ("FX Rate Source & Convention", "OANDA spot, period-end"),
    ]
    for i, (field_name, value) in enumerate(fields, start=1):
        ws.cell(row=header_excel_row + i, column=2, value=field_name)
        ws.cell(row=header_excel_row + i, column=3, value=value)

    # Bank tab: renamed sheet, header one row lower than canonical, renamed columns.
    ws = wb.create_sheet("Bank Confo Export")
    bank_header_row = config.POPULATION_HEADER_ROW + 1 + 1  # 0-indexed override -> 1-indexed Excel row
    bank_columns = [
        "Entity/Component Name",
        "Jurisdiction / Template Standard",
        "Bank",  # renamed from "Bank Name"
        "Contact Email",  # renamed from "Bank Contact Email"
        "Acct #",  # renamed from "Account Number(s)"
        "GL Balance (Local Currency)",
        "Local Currency",
        "Confirmation Form Type",
    ]
    bank_rows = [
        ["Acme US Inc.", "US / ASB 92", "First National Bank", "confirms@fnb.example.com", "AC-1001", 250000, "USD", "Standard"],
    ]
    _write_table(ws, bank_header_row, bank_columns, bank_rows)

    # Everything else already matches io-spec.md canonically — headers only, no data rows needed.
    for sheet_name, columns in (
        (config.SHEET_AR, schemas.AR_COLUMNS),
        (config.SHEET_LEGAL, schemas.LEGAL_COLUMNS),
        (config.SHEET_INTERCOMPANY, schemas.INTERCOMPANY_COLUMNS),
        (config.SHEET_GL_EXTRACT, schemas.GL_COLUMNS),
        (config.SHEET_RELATED_PARTY, schemas.RELATED_PARTY_COLUMNS),
    ):
        ws = wb.create_sheet(sheet_name)
        _write_table(ws, config.POPULATION_HEADER_ROW + 1, columns, [])

    out_path = tmp_path / "client_export.xlsx"
    wb.save(out_path)
    return out_path


def test_mismatched_layout_fails_without_adapter(tmp_path):
    path = _build_client_workbook(tmp_path)
    with pytest.raises(ValueError):
        load_all(path)


def test_adapter_normalizes_mismatched_layout(tmp_path):
    path = _build_client_workbook(tmp_path)
    result = load_all(path, adapter="example_client")

    assert result.engagement["Client Name"] == "Acme Test Co"
    assert result.engagement["Engagement Code"] == "ACM-FY26-001"

    assert len(result.bank) == 1
    row = result.bank.iloc[0]
    assert row["Bank Name"] == "First National Bank"
    assert row["Bank Contact Email"] == "confirms@fnb.example.com"
    assert row["Account Number(s)"] == "AC-1001"

    # No spurious Tier 1 gates from columns the adapter forgot to map
    # (ValidationIssue.sheet is always the canonical sheet name, per intake.py).
    bank_issues = [i for i in result.issues if i.sheet == config.SHEET_BANK]
    assert bank_issues == []


def test_unknown_adapter_name_raises_helpful_error():
    with pytest.raises(KeyError):
        get_adapter("does_not_exist")
